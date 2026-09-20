from __future__ import annotations

import hashlib
import json
import os
import platform
import sys
import time
from pathlib import Path

os.environ.setdefault("OMP_NUM_THREADS", "1")
os.environ.setdefault("MKL_NUM_THREADS", "1")
os.environ.setdefault("OPENBLAS_NUM_THREADS", "1")
os.environ.setdefault("NUMEXPR_NUM_THREADS", "1")

import numpy as np
import pandas as pd
from openpyxl import load_workbook


PROJECT = Path(__file__).resolve().parents[1]
RAW_EXISTING = PROJECT / "data" / "raw" / "nhanes_2021_2023"
OUTPUT = PROJECT / "results" / "schema_expansion"

CONFIG_XLSX = PROJECT / "configs" / "schema_expansion_design.xlsx"
METADATA_CSV = PROJECT / "data" / "processed" / "working" / "nhanes_metadata.csv"

EPSILONS = [0.1, 0.25, 0.5, 1.0, 2.0, 4.0, 8.0]
MECHANISMS = ["GRR", "OUE"]
RUNS = 30
RANDOM_PARTITIONS = 100
BASE_SEED = 2026
DOMAIN_SIZE = 3

sys.path.insert(0, str(PROJECT))
from semmed_ldp.ldp_mechanisms import (  # noqa: E402
    grr_estimate,
    grr_perturb,
    oue_estimate,
    oue_perturb,
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def worksheet_frame(path: Path, sheet: str) -> pd.DataFrame:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet]
    values = list(worksheet.iter_rows(values_only=True))
    return pd.DataFrame(values[1:], columns=values[0])


def source_path(component: str) -> Path:
    candidate = RAW_EXISTING / f"{component}.XPT"
    if candidate.exists():
        return candidate
    raise FileNotFoundError(f"Missing NHANES source component: {component}.XPT")


def load_design() -> tuple[pd.DataFrame, pd.DataFrame]:
    attributes = worksheet_frame(CONFIG_XLSX, "Attribute_Config")
    rules = worksheet_frame(CONFIG_XLSX, "Rule_Lookup")
    attributes = attributes.loc[attributes["ID"].notna()].copy()
    rules = rules.loc[rules["Attribute ID"].notna()].copy()
    attributes["ID"] = attributes["ID"].astype(int)
    rules["Attribute ID"] = rules["Attribute ID"].astype(int)
    if attributes["ID"].tolist() != list(range(1, 41)):
        raise ValueError("The frozen design must contain attribute IDs 1..40 in order")
    if not attributes["Freeze status"].eq("FROZEN").all():
        raise ValueError("Every extension attribute must be frozen before analysis")
    return attributes, rules


def load_raw_panel(attributes: pd.DataFrame, metadata: pd.DataFrame) -> tuple[pd.DataFrame, list[Path]]:
    components: dict[str, pd.DataFrame] = {}
    source_files: list[Path] = []
    for component in attributes["NHANES file"].drop_duplicates():
        path = source_path(str(component))
        source_files.append(path)
        needed = ["SEQN"] + attributes.loc[
            attributes["NHANES file"].eq(component), "NHANES variable"
        ].astype(str).tolist()
        frame = pd.read_sas(path, format="xport", encoding="utf-8")
        missing = sorted(set(needed).difference(frame.columns))
        if missing:
            raise ValueError(f"{path.name} lacks required variables: {missing}")
        components[str(component)] = frame[needed].copy()

    merged = metadata[["SEQN", "age_years", "sex"]].copy()
    for component in attributes["NHANES file"].drop_duplicates():
        merged = merged.merge(
            components[str(component)], on="SEQN", how="left", validate="one_to_one"
        )
    if len(merged) != len(metadata) or merged["SEQN"].duplicated().any():
        raise ValueError("The fixed adult cohort changed during component merging")
    return merged, source_files


def apply_rules(
    merged: pd.DataFrame, attributes: pd.DataFrame, rules: pd.DataFrame
) -> tuple[pd.DataFrame, pd.DataFrame]:
    states = pd.DataFrame(index=merged.index)
    coverage_rows: list[dict[str, object]] = []
    sex = merged["sex"].astype(str).str.lower()
    age = pd.to_numeric(merged["age_years"], errors="coerce")

    for _, attribute in attributes.iterrows():
        attribute_id = int(attribute["ID"])
        variable = str(attribute["NHANES variable"])
        values = pd.to_numeric(merged[variable], errors="coerce")
        applicable = rules.loc[rules["Attribute ID"].eq(attribute_id)]
        matched = np.zeros(len(merged), dtype=int)
        normal = np.zeros(len(merged), dtype=bool)

        for _, rule in applicable.iterrows():
            mask = np.ones(len(merged), dtype=bool)
            rule_sex = str(rule["Sex"]).strip().lower()
            if rule_sex != "any":
                mask &= sex.eq(rule_sex).to_numpy()
            age_min = rule["Age min (yr, inclusive)"]
            age_max = rule["Age max (yr, exclusive)"]
            if pd.notna(age_min):
                mask &= age.ge(float(age_min)).to_numpy()
            if pd.notna(age_max):
                mask &= age.lt(float(age_max)).to_numpy()
            matched += mask.astype(int)

            local_normal = mask.copy()
            lower = rule["Lower"]
            upper = rule["Upper"]
            if pd.notna(lower):
                if bool(rule["Lower inclusive"]):
                    local_normal &= values.ge(float(lower)).to_numpy()
                else:
                    local_normal &= values.gt(float(lower)).to_numpy()
            if pd.notna(upper):
                if bool(rule["Upper inclusive"]):
                    local_normal &= values.le(float(upper)).to_numpy()
                else:
                    local_normal &= values.lt(float(upper)).to_numpy()
            normal |= local_normal

        nonmissing = values.notna().to_numpy()
        bad_coverage = nonmissing & (matched != 1)
        if bad_coverage.any():
            raise ValueError(
                f"Rule coverage failure for {variable}: {int(bad_coverage.sum())} nonmissing records"
            )
        encoded = np.full(len(merged), 2, dtype=np.int8)
        encoded[nonmissing] = 1
        encoded[nonmissing & normal] = 0
        states[variable] = encoded
        coverage_rows.append(
            {
                "attribute_id": attribute_id,
                "attribute": attribute["Attribute"],
                "nhanes_variable": variable,
                "n_total": len(merged),
                "n_nonmissing": int(nonmissing.sum()),
                "n_missing": int((~nonmissing).sum()),
                "missing_rate": float((~nonmissing).mean()),
                "n_normal": int((encoded == 0).sum()),
                "n_abnormal": int((encoded == 1).sum()),
                "normal_rate_among_nonmissing": float((encoded[nonmissing] == 0).mean()),
                "abnormal_rate_among_nonmissing": float((encoded[nonmissing] == 1).mean()),
                "rule_rows": int(len(applicable)),
            }
        )
    return states, pd.DataFrame(coverage_rows)


def aggregate_groups(attribute_states: np.ndarray, labels: list[str]) -> tuple[np.ndarray, list[str]]:
    label_order = list(dict.fromkeys(labels))
    result = np.empty((attribute_states.shape[0], len(label_order)), dtype=np.int8)
    label_array = np.asarray(labels, dtype=object)
    for j, label in enumerate(label_order):
        local = attribute_states[:, label_array == label]
        result[:, j] = 0
        result[np.all(local == 2, axis=1), j] = 2
        result[np.any(local == 1, axis=1), j] = 1
    return result, label_order


def random_labels(group_sizes: list[int], n_attributes: int, seed: int) -> list[str]:
    labels = [f"random_{i + 1}" for i, size in enumerate(group_sizes) for _ in range(size)]
    if len(labels) != n_attributes:
        raise ValueError("Random partition size profile does not match the panel")
    rng = np.random.default_rng(seed)
    rng.shuffle(labels)
    return labels


def truth_distribution(states: np.ndarray) -> np.ndarray:
    n, m = states.shape
    truth = np.empty((m, DOMAIN_SIZE), dtype=float)
    for j in range(m):
        truth[j] = np.bincount(states[:, j], minlength=DOMAIN_SIZE) / n
    return truth


def smp_client(states: np.ndarray, epsilon: float, mechanism: str, seed: int):
    n, m = states.shape
    rng = np.random.default_rng(seed)
    selected = rng.integers(0, m, size=n)
    values = states[np.arange(n), selected]
    if mechanism == "GRR":
        reports = grr_perturb(values, epsilon, DOMAIN_SIZE, seed=seed + 1)
    else:
        reports = oue_perturb(values, epsilon, DOMAIN_SIZE, seed=seed + 1)
    return selected, reports


def smp_server(selected: np.ndarray, reports: np.ndarray, m: int, epsilon: float, mechanism: str):
    estimates = np.empty((m, DOMAIN_SIZE), dtype=float)
    report_counts = np.empty(m, dtype=int)
    for j in range(m):
        local = reports[selected == j]
        report_counts[j] = len(local)
        if len(local) == 0:
            estimates[j] = 1.0 / DOMAIN_SIZE
        elif mechanism == "GRR":
            estimates[j] = grr_estimate(local, epsilon, DOMAIN_SIZE)
        else:
            estimates[j] = oue_estimate(local, epsilon, DOMAIN_SIZE)
    return estimates, report_counts


def mse_for(states: np.ndarray, epsilon: float, mechanism: str, seed: int) -> tuple[float, np.ndarray]:
    truth = truth_distribution(states)
    selected, reports = smp_client(states, epsilon, mechanism, seed)
    estimate, report_counts = smp_server(selected, reports, states.shape[1], epsilon, mechanism)
    return float(np.mean((truth - estimate) ** 2)), report_counts


def state_distribution_rows(
    panel: str, method: str, states: np.ndarray, dimension_names: list[str], group_sizes: list[int]
) -> list[dict[str, object]]:
    truth = truth_distribution(states)
    rows: list[dict[str, object]] = []
    for j, name in enumerate(dimension_names):
        for state in range(DOMAIN_SIZE):
            rows.append(
                {
                    "panel": panel,
                    "method": method,
                    "dimension": name,
                    "group_size": int(group_sizes[j]),
                    "state": state,
                    "state_label": ["no_abnormality", "abnormality_present", "all_missing"][state],
                    "proportion": float(truth[j, state]),
                }
            )
    return rows


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    attributes, rules = load_design()
    metadata = pd.read_csv(METADATA_CSV)
    if len(metadata) != 8153 or metadata["age_years"].min() < 18:
        raise ValueError("Expected the frozen 8,153-person adult NHANES cohort")

    merged, source_files = load_raw_panel(attributes, metadata)
    attribute_states, attribute_summary = apply_rules(merged, attributes, rules)
    attribute_summary.to_csv(OUTPUT / "attribute_missingness_and_states.csv", index=False)

    # Public, non-patient-level configuration export for Supplementary Table S1.
    supplement_columns = [
        "ID", "Attribute", "NHANES file", "NHANES variable", "Native unit",
        "LOINC (NHANES)", "HPO ID", "HPO official anchor", "Semantic group", "Added in",
        "CDC source", "LOINC source", "HPO source",
    ]
    attributes[supplement_columns].to_csv(OUTPUT / "supplement_table_s1_mapping.csv", index=False)
    rules.to_csv(OUTPUT / "frozen_rule_lookup.csv", index=False)
    adult_rules = rules.loc[
        rules["Age max (yr, exclusive)"].isna()
        | pd.to_numeric(rules["Age max (yr, exclusive)"], errors="coerce").gt(18)
    ].copy()
    adult_rules["rule_display"] = adult_rules.apply(
        lambda row: (
            f"{row['Sex']}, age >= {max(18, int(row['Age min (yr, inclusive)']))}: "
            f"{row['Normal rule']}"
        ),
        axis=1,
    )
    rule_summary = (
        adult_rules.groupby("Attribute ID", sort=False)["rule_display"]
        .apply("; ".join)
        .to_dict()
    )
    supplement_s1 = pd.DataFrame(
        {
            "Attribute": attributes["Attribute"],
            "NHANES Variable": attributes["NHANES variable"],
            "LOINC": attributes["LOINC (NHANES)"],
            "Unit": attributes["Native unit"],
            "Reference/Decision Rule": attributes["ID"].map(rule_summary),
            "HPO Anchor": attributes["HPO ID"] + " (" + attributes["HPO official anchor"] + ")",
            "Final Semantic Group": attributes["Semantic group"],
            "Panel": attributes["Added in"],
        }
    )
    supplement_s1.to_csv(OUTPUT / "supplement_table_s1.csv", index=False)

    panel_rows: list[dict[str, object]] = []
    partition_rows: list[dict[str, object]] = []
    raw_metric_rows: list[dict[str, object]] = []
    collapsed_metric_rows: list[dict[str, object]] = []
    distribution_rows: list[dict[str, object]] = []
    runtime_rows: list[dict[str, object]] = []
    for d in [10, 20, 40]:
        panel = f"D{d}"
        panel_attributes = attributes.iloc[:d].copy()
        columns = panel_attributes["NHANES variable"].astype(str).tolist()
        raw_states = attribute_states[columns].to_numpy(dtype=np.int8)
        semantic_labels = panel_attributes["Semantic group"].astype(str).tolist()
        kgr_states, group_names = aggregate_groups(raw_states, semantic_labels)
        group_sizes = [semantic_labels.count(name) for name in group_names]
        m = len(group_names)
        panel_rows.append(
            {
                "panel": panel,
                "n_users": len(raw_states),
                "n_attributes": d,
                "n_semantic_dimensions": m,
                "compression_ratio_d_over_m": d / m,
                "mean_group_size": float(np.mean(group_sizes)),
                "max_group_size": int(max(group_sizes)),
                "expected_reports_per_kgr_dimension": len(raw_states) / m,
                "expected_reports_per_raw_dimension": len(raw_states) / d,
            }
        )
        distribution_rows.extend(
            state_distribution_rows(panel, "KGR-LDP", kgr_states, group_names, group_sizes)
        )
        distribution_rows.extend(
            state_distribution_rows(panel, "Raw", raw_states, columns, [1] * d)
        )

        random_states: list[np.ndarray] = []
        for partition in range(RANDOM_PARTITIONS):
            partition_seed = BASE_SEED + 1000 + partition
            labels = random_labels(group_sizes, d, partition_seed)
            states, random_names = aggregate_groups(raw_states, labels)
            random_states.append(states)
            for attribute_id, variable, group in zip(
                panel_attributes["ID"].astype(int), columns, labels
            ):
                partition_rows.append(
                    {
                        "panel": panel,
                        "partition": partition + 1,
                        "seed": partition_seed,
                        "attribute_id": attribute_id,
                        "nhanes_variable": variable,
                        "random_group": group,
                    }
                )

        for run in range(RUNS):
            run_seed = BASE_SEED + run
            # Match the formal-v3 full-cohort path exactly: selecting all users
            # still returns a seeded permutation before local reporting.
            order = np.random.default_rng(run_seed).choice(
                len(raw_states), size=len(raw_states), replace=False
            )
            for mechanism in MECHANISMS:
                for epsilon in EPSILONS:
                    for method, states in [("KGR-LDP", kgr_states), ("Raw", raw_states)]:
                        ordered_states = states[order]
                        mse, report_counts = mse_for(
                            ordered_states, epsilon, mechanism, run_seed
                        )
                        row = {
                            "panel": panel,
                            "method": method,
                            "mechanism": mechanism,
                            "epsilon": epsilon,
                            "run": run,
                            "seed": run_seed,
                            "n_users": len(states),
                            "n_dimensions": states.shape[1],
                            "MSE": mse,
                            "min_reports_per_dimension": int(report_counts.min()),
                            "max_reports_per_dimension": int(report_counts.max()),
                        }
                        raw_metric_rows.append(row)
                        collapsed_metric_rows.append(row)

                    partition_mse: list[float] = []
                    min_reports: list[int] = []
                    max_reports: list[int] = []
                    for partition, states in enumerate(random_states, start=1):
                        ordered_states = states[order]
                        mse, report_counts = mse_for(
                            ordered_states, epsilon, mechanism, run_seed
                        )
                        partition_mse.append(mse)
                        min_reports.append(int(report_counts.min()))
                        max_reports.append(int(report_counts.max()))
                        raw_metric_rows.append(
                            {
                                "panel": panel,
                                "method": "Random",
                                "partition": partition,
                                "mechanism": mechanism,
                                "epsilon": epsilon,
                                "run": run,
                                "seed": run_seed,
                                "n_users": len(states),
                                "n_dimensions": states.shape[1],
                                "MSE": mse,
                                "min_reports_per_dimension": int(report_counts.min()),
                                "max_reports_per_dimension": int(report_counts.max()),
                            }
                        )
                    collapsed_metric_rows.append(
                        {
                            "panel": panel,
                            "method": "Random",
                            "mechanism": mechanism,
                            "epsilon": epsilon,
                            "run": run,
                            "seed": run_seed,
                            "n_users": len(kgr_states),
                            "n_dimensions": m,
                            "MSE": float(np.mean(partition_mse)),
                            "random_partition_sd": float(np.std(partition_mse, ddof=1)),
                            "min_reports_per_dimension": int(min(min_reports)),
                            "max_reports_per_dimension": int(max(max_reports)),
                        }
                    )

        # Runtime: KGR-LDP only, epsilon=1, one warm-up followed by ten timed repeats.
        for mechanism in MECHANISMS:
            for repeat in range(11):
                seed = BASE_SEED + 9000 + repeat
                started = time.perf_counter()
                timed_attribute_states, _ = apply_rules(
                    merged, panel_attributes, rules
                )
                transformed, _ = aggregate_groups(
                    timed_attribute_states[columns].to_numpy(dtype=np.int8),
                    semantic_labels,
                )
                selected, reports = smp_client(transformed, 1.0, mechanism, seed)
                client_seconds = time.perf_counter() - started
                started = time.perf_counter()
                _, report_counts = smp_server(selected, reports, transformed.shape[1], 1.0, mechanism)
                server_seconds = time.perf_counter() - started
                if repeat > 0:
                    runtime_rows.append(
                        {
                            "panel": panel,
                            "mechanism": mechanism,
                            "epsilon": 1.0,
                            "repeat": repeat,
                            "n_users": len(transformed),
                            "n_dimensions": transformed.shape[1],
                            "client_seconds": client_seconds,
                            "server_seconds": server_seconds,
                            "total_seconds": client_seconds + server_seconds,
                            "min_reports_per_dimension": int(report_counts.min()),
                            "max_reports_per_dimension": int(report_counts.max()),
                        }
                    )

    panel_summary = pd.DataFrame(panel_rows)
    partitions = pd.DataFrame(partition_rows)
    raw_metrics = pd.DataFrame(raw_metric_rows)
    run_metrics = pd.DataFrame(collapsed_metric_rows)
    state_distributions = pd.DataFrame(distribution_rows)
    runtime_runs = pd.DataFrame(runtime_rows)

    summary = (
        run_metrics.groupby(["panel", "method", "mechanism", "epsilon"], as_index=False)
        .agg(
            mean_MSE=("MSE", "mean"),
            sd_MSE=("MSE", "std"),
            median_MSE=("MSE", "median"),
            min_MSE=("MSE", "min"),
            max_MSE=("MSE", "max"),
            n_runs=("MSE", "size"),
        )
    )
    core = summary.loc[summary["epsilon"].eq(1.0)].copy()
    runtime_summary = (
        runtime_runs.groupby(["panel", "mechanism", "epsilon"], as_index=False)
        .agg(
            median_client_seconds=("client_seconds", "median"),
            median_server_seconds=("server_seconds", "median"),
            median_total_seconds=("total_seconds", "median"),
            n_timed_repeats=("repeat", "size"),
        )
    )

    panel_summary.to_csv(OUTPUT / "panel_summary.csv", index=False)
    partitions.to_csv(OUTPUT / "random_partition_assignments.csv", index=False)
    raw_metrics.to_csv(OUTPUT / "all_partition_run_metrics.csv", index=False)
    run_metrics.to_csv(OUTPUT / "run_metrics_random_collapsed.csv", index=False)
    summary.to_csv(OUTPUT / "mse_summary.csv", index=False)
    core.to_csv(OUTPUT / "core_epsilon1_table.csv", index=False)
    state_distributions.to_csv(OUTPUT / "state_concentration.csv", index=False)
    runtime_runs.to_csv(OUTPUT / "runtime_repeats.csv", index=False)
    runtime_summary.to_csv(OUTPUT / "runtime_summary.csv", index=False)

    source_records = [
        {
            "path": path.relative_to(PROJECT).as_posix(),
            "bytes": path.stat().st_size,
            "sha256": sha256(path),
        }
        for path in [CONFIG_XLSX, METADATA_CSV, *source_files]
    ]
    manifest = {
        "experiment": "Experiment 5: dimensional scaling D10/D20/D40",
        "status": "complete",
        "generated_at_local": time.strftime("%Y-%m-%d %H:%M:%S"),
        "cohort": "NHANES August 2021-August 2023, adults aged >=18",
        "n_users": int(len(metadata)),
        "panels": [10, 20, 40],
        "epsilon_list": EPSILONS,
        "mechanisms": MECHANISMS,
        "strategy": "Smp",
        "repeat_runs": RUNS,
        "random_partitions": RANDOM_PARTITIONS,
        "base_seed": BASE_SEED,
        "state_encoding": {"0": "no abnormality", "1": "abnormality present", "2": "all values missing"},
        "random_aggregation": "MSE averaged over 100 fixed size-matched partitions within each perturbation run",
        "runtime_protocol": "KGR-LDP; one warm-up plus ten timed repeats; vectorized NumPy implementation",
        "software": {
            "python": sys.version,
            "platform": platform.platform(),
            "processor": platform.processor(),
            "logical_cpu_count": os.cpu_count(),
            "numpy": np.__version__,
            "pandas": pd.__version__,
            "multithreading": "disabled through OMP/MKL/OpenBLAS/NumExpr environment variables",
        },
        "inputs": source_records,
        "outputs": sorted(path.name for path in OUTPUT.glob("*.csv")),
    }
    (OUTPUT / "manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8"
    )

    print(panel_summary.to_string(index=False))
    print("\nCore epsilon=1 results:\n", core.to_string(index=False))
    print("\nRuntime summary:\n", runtime_summary.to_string(index=False))
    print(f"\nWrote outputs to {OUTPUT}")


if __name__ == "__main__":
    main()
